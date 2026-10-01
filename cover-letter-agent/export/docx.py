"""Exporter: approved letter -> .docx / .pdf in output/.

Only approved letters export, and the file always holds the version you approved.
The template is simple: your name and contact details at the top, the date, then
the letter. PDF goes through docx2pdf (Microsoft Word) or LibreOffice.
"""

import io
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

import config

EXPORTABLE = ("approved", "submitted")
FONT = "Calibri"
PDF_TIMEOUT = 120


class NotApproved(Exception):
    """Final export is only for letters you've approved."""


class PdfUnavailable(Exception):
    """Neither Word nor LibreOffice could make the PDF."""


def can_export(app: dict) -> bool:
    return app.get("status") in EXPORTABLE and app.get("sent_version") is not None


def require_approved(app: dict) -> None:
    if not can_export(app):
        raise NotApproved("Approve the letter before exporting it. Drafts can be copied, not exported.")


def file_stem(app: dict) -> str:
    """'Company - Title - Cover Letter', with characters Windows won't allow removed."""
    parts = [app.get("company") or "Company", app.get("title") or "Job", "Cover Letter"]
    stem = " - ".join(re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", p).strip() for p in parts)
    return re.sub(r"\s+", " ", stem)[:150].rstrip(" .")


def _short_url(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url.strip()).rstrip("/")


def contact_line(profile) -> str:
    phone = f"{profile.country_code} {profile.phone}".strip() if profile.phone else ""
    links = [_short_url(u) for u in (profile.linkedin, profile.portfolio) if u]
    items = [profile.city, profile.email, phone, *links]
    return " · ".join(i.strip() for i in items if i and i.strip())


def _plain(text: str) -> str:
    """Drop Markdown emphasis the letter shouldn't have anyway."""
    return re.sub(r"(\*\*|__|\*|_)(\S.*?\S|\S)\1", r"\2", text)


def build_docx(letter: str, profile, when: date | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(72)
        section.left_margin = section.right_margin = Pt(72)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    if profile.name:
        head = doc.add_paragraph()
        run = head.add_run(profile.name)
        run.bold = True
        run.font.size = Pt(16)
        head.paragraph_format.space_after = Pt(2)
    contact = contact_line(profile)
    if contact:
        line = doc.add_paragraph()
        run = line.add_run(contact)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x4F, 0x5B, 0x66)
        line.paragraph_format.space_after = Pt(18)

    when = when or date.today()
    doc.add_paragraph(f"{when:%B} {when.day}, {when.year}")

    for block in _plain(letter).strip().split("\n\n"):
        paragraph = doc.add_paragraph()
        lines = block.split("\n")
        for i, line in enumerate(lines):
            run = paragraph.add_run(line.strip())
            if i < len(lines) - 1:
                run.add_break()  # sign-off and name stay on separate lines

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _approved_text(app: dict, drafts: list[dict]) -> str:
    version = next((d for d in drafts if d["version"] == app["sent_version"]), None)
    if version is None:
        raise NotApproved("The approved version couldn't be found.")
    return version["text"]


def export_docx(app: dict, drafts: list[dict], profile, folder: Path | None = None) -> Path:
    require_approved(app)
    folder = folder or config.OUTPUT_DIR
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{file_stem(app)}.docx"
    path.write_bytes(build_docx(_approved_text(app, drafts), profile))
    return path


def _soffice() -> str | None:
    found = shutil.which("soffice")
    if found:
        return found
    for candidate in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                      "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(candidate).exists():
            return candidate
    return None


def docx_to_pdf(docx_path: Path) -> Path:
    """Convert next to the .docx. Runs in a separate process so a stuck Word can be stopped."""
    pdf_path = docx_path.with_suffix(".pdf")
    pdf_path.unlink(missing_ok=True)
    errors = []
    try:
        subprocess.run(
            [sys.executable, "-c", "import sys; from docx2pdf import convert; convert(sys.argv[1], sys.argv[2])",
             str(docx_path), str(pdf_path)],
            check=True, capture_output=True, timeout=PDF_TIMEOUT,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        errors.append(f"Word: {exc}")
    if not pdf_path.exists() and (soffice := _soffice()):
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                            str(docx_path)], check=True, capture_output=True, timeout=PDF_TIMEOUT)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            errors.append(f"LibreOffice: {exc}")
    if not pdf_path.exists():
        raise PdfUnavailable(
            "Couldn't make the PDF. It needs Microsoft Word or LibreOffice installed. "
            "The .docx was saved and can be opened and saved as PDF by hand."
        )
    return pdf_path


def export_pdf(app: dict, drafts: list[dict], profile, folder: Path | None = None) -> Path:
    return docx_to_pdf(export_docx(app, drafts, profile, folder))
